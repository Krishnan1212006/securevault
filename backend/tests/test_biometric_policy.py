from __future__ import annotations

from pathlib import Path

from fastapi import HTTPException
import pytest

from backend import audit_service, biometric_auth, main
from backend import windows_biometrics
from backend.windows_biometrics import BiometricProviderError, WindowsBiometricFramework


@pytest.fixture
def isolated_audit(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    state = tmp_path / "state"
    monkeypatch.setattr(audit_service, "APP_DIR", state)
    monkeypatch.setattr(audit_service, "DB_PATH", state / "audit.sqlite3")
    monkeypatch.setattr(audit_service, "KEY_PATH", state / "audit.key")
    monkeypatch.setattr(main, "ENCRYPTED", {"selected": {
        "path": tmp_path / "private.svlt",
        "reference": None,
    }})


def detected_capabilities(fingerprint: dict | None = None, face: dict | None = None) -> dict:
    return {
        "operating_system": {"name": "Windows", "version": "test"},
        "provider_status": "AVAILABLE",
        "provider_message": "Mocked provider for unit test.",
        "provider": "Mock WBF",
        "factors": {
            "face": face or {
                "available": False, "status": "UNSUPPORTED", "provider_status": "PROVIDER_NOT_IMPLEMENTED",
                "hardware_status": "NOT_EXPOSED_BY_SUPPORTED_API", "enrollment_status": "UNKNOWN",
                "device_count": 0, "modality_identifiable": False,
                "message": "Face matching is not exposed by the supported API.",
            },
            "fingerprint": fingerprint or {
                "available": True, "status": "READY", "provider_status": "AVAILABLE",
                "hardware_status": "DETECTED", "enrollment_status": "ENROLLED",
                "device_count": 1, "modality_identifiable": True,
                "message": "Mock sensor and enrollment detected.",
            },
        },
        "two_factor_supported": False,
        "decryption_authorized": False,
    }


class MockProvider:
    def __init__(self, *, result: str = "success", detection: dict | None = None):
        self.result = result
        self.detection = detection or detected_capabilities()
        self.verification_calls = 0

    def detect(self) -> dict:
        return self.detection

    def verify_current_user_fingerprint(self) -> None:
        self.verification_calls += 1
        if self.result != "success":
            raise BiometricProviderError("Mock fingerprint verification failed.")


def test_capabilities_show_fingerprint_but_deny_two_factor_policy(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(biometric_auth, "provider", lambda: MockProvider())
    result = biometric_auth.capabilities()

    assert result["provider"] == "Mock WBF"
    assert result["decryption_authorized"] is False
    assert result["factors"]["face"]["status"] == "UNSUPPORTED"
    assert result["factors"]["face"]["modality_identifiable"] is False
    assert result["factors"]["fingerprint"]["available"] is True
    assert result["factors"]["fingerprint"]["status"] == "READY"
    assert result["two_factor_supported"] is False


@pytest.mark.parametrize("status", ["HARDWARE_NOT_DETECTED", "ENROLLMENT_MISSING"])
def test_fingerprint_capability_reports_missing_prerequisite(
    monkeypatch: pytest.MonkeyPatch, status: str,
) -> None:
    factor = {
        "available": False, "status": status, "provider_status": status,
        "hardware_status": "NOT_DETECTED" if status == "HARDWARE_NOT_DETECTED" else "DETECTED",
        "enrollment_status": "MISSING",
        "device_count": 0 if status == "HARDWARE_NOT_DETECTED" else 1,
        "modality_identifiable": True, "message": status,
    }
    monkeypatch.setattr(
        biometric_auth, "provider", lambda: MockProvider(detection=detected_capabilities(fingerprint=factor)),
    )
    assert biometric_auth.capabilities()["factors"]["fingerprint"]["status"] == status


@pytest.mark.parametrize(
    ("sensor_count", "enrolled_factors", "expected"),
    [
        (0, windows_biometrics.FINGERPRINT, "HARDWARE_NOT_DETECTED"),
        (1, 0, "ENROLLMENT_MISSING"),
        (1, windows_biometrics.FINGERPRINT, "READY"),
    ],
)
def test_windows_provider_detection_uses_sensor_and_current_account_enrollment(
    monkeypatch: pytest.MonkeyPatch, sensor_count: int, enrolled_factors: int, expected: str,
) -> None:
    provider = WindowsBiometricFramework()
    provider._dll = object()
    monkeypatch.setattr(provider, "_identity", lambda: (object(), object()))
    monkeypatch.setattr(provider, "_service_status", lambda: ("RUNNING", "Service is running."))
    monkeypatch.setattr(provider, "_enrolled_factors", lambda _identity: (enrolled_factors, None))
    monkeypatch.setattr(provider, "_enumerate_fingerprints", lambda: (sensor_count, None))
    monkeypatch.setattr(
        windows_biometrics.ctypes, "WinDLL",
        lambda *_args, **_kwargs: type("Kernel", (), {"CloseHandle": staticmethod(lambda _token: True)})(),
    )

    result = provider.detect()

    assert result["factors"]["fingerprint"]["status"] == expected
    assert result["factors"]["fingerprint"]["available"] is (expected == "READY")
    assert result["factors"]["face"]["status"] == "UNSUPPORTED"
    assert result["two_factor_supported"] is False


def test_face_enrollment_bit_does_not_claim_face_sensor_or_authentication(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    provider = WindowsBiometricFramework()
    provider._dll = object()
    monkeypatch.setattr(provider, "_identity", lambda: (object(), object()))
    monkeypatch.setattr(provider, "_service_status", lambda: ("RUNNING", "Service is running."))
    monkeypatch.setattr(
        provider, "_enrolled_factors",
        lambda _identity: (windows_biometrics.FACE | windows_biometrics.FINGERPRINT, None),
    )
    monkeypatch.setattr(provider, "_enumerate_fingerprints", lambda: (1, None))
    monkeypatch.setattr(
        windows_biometrics.ctypes, "WinDLL",
        lambda *_args, **_kwargs: type("Kernel", (), {"CloseHandle": staticmethod(lambda _token: True)})(),
    )

    face = provider.detect()["factors"]["face"]

    assert face["enrollment_status"] == "ENROLLED"
    assert face["hardware_status"] == "NOT_EXPOSED_BY_SUPPORTED_API"
    assert face["status"] == "UNSUPPORTED"
    assert face["available"] is False


def test_stopped_windows_biometric_service_is_not_reported_as_missing_hardware(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    provider = WindowsBiometricFramework()
    provider._dll = object()
    monkeypatch.setattr(provider, "_identity", lambda: (object(), object()))
    monkeypatch.setattr(provider, "_service_status", lambda: ("STOPPED", "Windows Biometric Service is stopped."))
    monkeypatch.setattr(provider, "_enrolled_factors", lambda _identity: (windows_biometrics.FINGERPRINT, None))
    monkeypatch.setattr(
        provider, "_enumerate_fingerprints",
        lambda: pytest.fail("Sensor discovery must not run while the biometric service is stopped."),
    )
    monkeypatch.setattr(
        windows_biometrics, "_close_handle", lambda _token: None,
    )

    result = provider.detect()

    assert result["windows_biometric_service"]["status"] == "STOPPED"
    assert result["native_provider_status"] == "UNAVAILABLE"
    assert result["factors"]["fingerprint"]["status"] == "PROVIDER_UNAVAILABLE"
    assert result["factors"]["fingerprint"]["hardware_status"] == "UNKNOWN"


@pytest.mark.parametrize(
    ("native_result", "expected_error"),
    [
        (0, None),
        (windows_biometrics.WINBIO_E_NO_MATCH, "did not match"),
        (windows_biometrics.WINBIO_E_CANCELED, "cancelled"),
        (windows_biometrics.HRESULT_ERROR_BUSY, "busy"),
    ],
)
def test_native_verify_uses_default_session_and_surfaces_native_results(
    monkeypatch: pytest.MonkeyPatch, native_result: int, expected_error: str | None,
) -> None:
    provider = WindowsBiometricFramework()
    identity = windows_biometrics._Identity()
    token = windows_biometrics.wintypes.HANDLE(1)
    opened: dict[str, int] = {}
    closed: list[int] = []

    def open_session(_factor, pool, flags, _units, unit_count, _database, session_out):
        opened.update(pool=pool, flags=flags, unit_count=unit_count)
        windows_biometrics.ctypes.cast(
            session_out, windows_biometrics.ctypes.POINTER(windows_biometrics.ctypes.c_void_p),
        )[0] = 99
        return 0

    def verify(_session, _identity, subfactor, _unit, match_out, _reject_detail):
        opened["subfactor"] = subfactor
        windows_biometrics.ctypes.cast(
            match_out, windows_biometrics.ctypes.POINTER(windows_biometrics.ctypes.c_ubyte),
        )[0] = native_result == 0
        return native_result

    def close_session(_session):
        closed.append(1)
        return 0

    winbio = type("WinBio", (), {
        "WinBioOpenSession": staticmethod(open_session),
        "WinBioVerify": staticmethod(verify),
        "WinBioCloseSession": staticmethod(close_session),
    })()
    provider._dll = winbio
    monkeypatch.setattr(provider, "detect", lambda: {
        "factors": {"fingerprint": {"available": True, "message": "Ready."}},
    })
    monkeypatch.setattr(provider, "_identity", lambda: (identity, token))
    monkeypatch.setattr(windows_biometrics, "_close_handle", lambda _token: None)

    if expected_error:
        with pytest.raises(BiometricProviderError, match=expected_error):
            provider.verify_current_user_fingerprint()
    else:
        provider.verify_current_user_fingerprint()

    assert opened == {
        "pool": windows_biometrics.WINBIO_POOL_SYSTEM,
        "flags": windows_biometrics.WINBIO_SESSION_FLAG_DEFAULT,
        "unit_count": 0,
        "subfactor": windows_biometrics.WINBIO_SUBTYPE_ANY,
    }
    assert closed == [1]


def test_authentication_is_not_attempted_without_ready_hardware_or_enrollment(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    fingerprint = {
        "available": False, "status": "ENROLLMENT_MISSING", "provider_status": "ENROLLMENT_MISSING",
        "hardware_status": "DETECTED", "enrollment_status": "MISSING", "device_count": 1,
        "modality_identifiable": True, "message": "Enroll a fingerprint first.",
    }
    provider = MockProvider(detection=detected_capabilities(fingerprint=fingerprint))
    monkeypatch.setattr(biometric_auth, "provider", lambda: provider)

    result = biometric_auth.authenticate("fingerprint")

    assert result["status"] == "ENROLLMENT_MISSING"
    assert result["authorization_id"] is None
    assert provider.verification_calls == 0


def test_cancelled_native_verification_is_not_reported_as_success(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    class CancelledProvider(MockProvider):
        def verify_current_user_fingerprint(self) -> None:
            raise windows_biometrics.BiometricAuthenticationCancelled("Windows prompt cancelled.")

    monkeypatch.setattr(biometric_auth.sys, "platform", "win32")
    monkeypatch.setattr(biometric_auth, "provider", lambda: CancelledProvider())

    result = biometric_auth.authenticate("fingerprint")

    assert result["status"] == "CANCELLED"
    assert result["authorization_id"] is None


@pytest.mark.parametrize(("factor", "expected"), [("face", "UNSUPPORTED"), ("fingerprint", "FAILED")])
def test_client_cannot_claim_biometric_success_and_attempt_is_audited(
    factor: str, expected: str, isolated_audit: None, monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(biometric_auth, "provider", lambda: MockProvider(result="failure"))
    body = main.BiometricAuthRequest.model_validate({
        "factor": factor,
        "encrypted_ids": ["selected"],
        "status": "SUCCESS",
        "verified": True,
    })

    result = main.authenticate_biometric(body, main.SESSION_TOKEN)

    assert result["status"] == expected
    assert result["authorization_id"] is None
    assert result["expires_at"] is None
    activity = audit_service.list_activities()[0]
    assert activity["operation"] == f"BIOMETRIC_{factor.upper()}"
    assert activity["status"] == expected
    assert activity["authentication_result"] == expected
    assert activity[f"{factor}_authentication_status"] == expected
    stats = audit_service.dashboard()
    assert stats["face_authentications"] == 0
    assert stats["fingerprint_authentications"] == 0
    assert stats["failed_authentication_attempts"] == (1 if factor == "fingerprint" else 0)


def test_success_is_reported_only_after_backend_provider_verifies(
    isolated_audit: None, monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(biometric_auth, "provider", lambda: MockProvider())
    body = main.BiometricAuthRequest(factor="fingerprint", encrypted_ids=["selected"])

    result = main.authenticate_biometric(body, main.SESSION_TOKEN)

    assert result["status"] == "SUCCESS"
    assert result["authorization_id"] is None
    assert audit_service.list_activities()[0]["authentication_result"] == "SUCCESS"
    assert audit_service.dashboard()["fingerprint_authentications"] == 1
    assert audit_service.dashboard()["face_authentications"] == 0


def test_mandatory_two_factor_decryption_is_denied_without_face_provider(
    isolated_audit: None, monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(biometric_auth, "provider", lambda: MockProvider())
    request = main.DecryptRequest(
        ids=["selected"], output_id="destination", password="test-passphrase",
        authentication_mode="face_and_fingerprint",
    )

    with pytest.raises(HTTPException) as error:
        main.start_decrypt(request, main.SESSION_TOKEN)

    assert error.value.status_code == 501
    record = audit_service.list_activities()[0]
    assert record["status"] == "FAILED"
    assert record["authentication_result"] == "REQUIRED_BIOMETRIC_FACTORS_UNAVAILABLE"
    assert record["face_authentication_status"] == "UNSUPPORTED"
    assert record["fingerprint_authentication_status"] == "READY"


def test_biometric_attempt_requires_backend_selected_encrypted_file(isolated_audit: None) -> None:
    body = main.BiometricAuthRequest(factor="face", encrypted_ids=["invented"])

    with pytest.raises(HTTPException) as error:
        main.authenticate_biometric(body, main.SESSION_TOKEN)

    assert error.value.status_code == 400
    assert audit_service.list_activities() == []


def test_audit_report_includes_biometric_and_decryption_fields(isolated_audit: None) -> None:
    from io import BytesIO

    from openpyxl import load_workbook

    audit_service.record_activity({
        "id": "auth", "filename": "private.svlt", "operation": "BIOMETRIC_FACE",
        "timestamp": "2026-01-01T00:00:00Z", "status": "UNSUPPORTED",
        "relative_source_path": "private.svlt", "file_size": 0,
        "face_authentication_status": "UNSUPPORTED", "authentication_result": "UNSUPPORTED",
        "error_details": "No separate face verification provider.",
    })
    workbook = load_workbook(BytesIO(main.make_report(audit_service.list_activities())))
    sheet = workbook.active
    headers = [cell.value for cell in sheet[1]]
    assert "Face authentication status" in headers
    assert "Fingerprint authentication status" in headers
    assert "Decrypted output path" in headers
    assert sheet.cell(row=2, column=headers.index("Face authentication status") + 1).value == "UNSUPPORTED"
