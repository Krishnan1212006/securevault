from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

HEADERS = [
    "Record ID", "Original filename", "Relative source path", "File size (bytes)", "Hash algorithm",
    "Original SHA-256 hash", "Encrypted output path", "Operation type", "Operation timestamp",
    "Face authentication status", "Fingerprint authentication status", "Decryption status",
    "Decrypted output path", "Recovered SHA-256 hash", "Authentication result",
    "Operation status", "Verification result", "Error details",
]


def _safe_cell(value):
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@", "\t", "\r")):
        return "'" + value
    return value


def make_report(records: list[dict]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "SecureVault Audit"
    sheet.append(HEADERS)
    for record in records:
        sheet.append([_safe_cell(value) for value in (
            record["id"], record["filename"], record["relative_source_path"], record["file_size"], "SHA-256",
            record.get("sha256") or "", record.get("encrypted_path") or "", record["operation"], record["timestamp"],
            record.get("face_authentication_status") or "", record.get("fingerprint_authentication_status") or "",
            record["status"] if record["operation"] == "DECRYPT" else "",
            record.get("decrypted_path") or "", record.get("recovered_sha256") or "",
            record.get("authentication_result") or "", record["status"],
            record.get("verification_result") or "", record.get("error_details") or "",
        )])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="10243A")
        cell.alignment = Alignment(horizontal="left")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column, width in {"A": 38, "B": 28, "C": 42, "D": 18, "E": 16, "F": 68,
                          "G": 48, "H": 18, "I": 28, "J": 22, "K": 24, "L": 18,
                          "M": 48, "N": 68, "O": 20, "P": 18, "Q": 22, "R": 48}.items():
        sheet.column_dimensions[column].width = width
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
